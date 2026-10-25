"""Uploaded documents for S0: text extraction, chunking, and one FAISS index per file.

Supported: PDF (pypdf), Word .docx (python-docx), PowerPoint .pptx (python-pptx), Excel .xlsx
(openpyxl), CSV/TSV, and plain text (.txt, .md, .json, .html, ...). Each file becomes units that
carry a location (page, slide, Word heading, sheet rows). Units are cut into overlapping,
sentence-aligned chunks sized for the MiniLM embedder (256 word pieces), embedded once with the
same model the stages use, and stored in a flat inner-product FAISS index: the vectors are
unit-normalised, so inner product is cosine and search is exact. Files are kept under
s0.store_dir as <id>/meta.json + <id>/index.faiss, keyed by content hash, so uploading the same
file twice reuses its index and the library survives a restart.
"""
from __future__ import annotations

import csv
import hashlib
import html
import io
import json
import logging
import os
import shutil
import threading
from dataclasses import dataclass
from datetime import datetime, timezone
from functools import lru_cache
from pathlib import Path

import numpy as np
import regex as re

from . import resources
from .config import thresholds
from .schemas import DocumentInfo

log = logging.getLogger(__name__)

BACKEND_DIR = Path(__file__).resolve().parents[1]
TEXT_EXTS = (".txt", ".md", ".markdown", ".rst", ".log", ".json", ".jsonl", ".yaml", ".yml", ".xml", ".html", ".htm")
LEGACY_EXTS = {".doc": ".docx", ".ppt": ".pptx", ".xls": ".xlsx"}


class UnsupportedFile(ValueError):
    """The file type is not handled, or the file holds no readable text."""


@dataclass
class Unit:
    text: str
    location: str | None = None


# ── Extraction ───────────────────────────────────────────────────────────────
def _decode(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def _pdf(data: bytes) -> tuple[list[Unit], int]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted and not reader.decrypt(""):
        raise UnsupportedFile("the PDF is password-protected")
    return [Unit(page.extract_text() or "", f"p. {i}") for i, page in enumerate(reader.pages, 1)], len(reader.pages)


def _row_text(cells) -> str:
    out: list[str] = []
    for c in cells:                      # merged cells repeat in python-docx / python-pptx rows
        c = " ".join(c.split())
        if c and (not out or out[-1] != c):
            out.append(c)
    return " | ".join(out)


def _docx(data: bytes) -> tuple[list[Unit], None]:
    import docx
    from docx.table import Table

    units: list[Unit] = []
    heading: str | None = None
    lines: list[str] = []
    has_body = False                     # headings with nothing under them yet join the next section

    for block in docx.Document(io.BytesIO(data)).iter_inner_content():
        if isinstance(block, Table):
            lines += [t for t in (_row_text(c.text for c in row.cells) for row in block.rows) if t]
            has_body = True
            continue
        text = block.text.strip()
        if not text:
            continue
        style = block.style.name if block.style is not None else ""
        if style.startswith(("Heading", "Title")):
            if has_body:
                units.append(Unit("\n".join(lines), heading))
                lines, has_body = [], False
            heading = text[:80]
        else:
            has_body = True
        lines.append(text)
    if lines:
        units.append(Unit("\n".join(lines), heading))
    return units, None


def _shape_lines(shape) -> list[str]:
    if hasattr(shape, "shapes"):                     # group shape
        return [line for s in shape.shapes for line in _shape_lines(s)]
    if getattr(shape, "has_table", False):
        return [t for t in (_row_text(c.text for c in row.cells) for row in shape.table.rows) if t]
    if shape.has_text_frame:
        return [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]
    return []


def _pptx(data: bytes) -> tuple[list[Unit], int]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    units = []
    for i, slide in enumerate(prs.slides, 1):
        lines = [line for shape in slide.shapes for line in _shape_lines(shape)]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"Speaker notes: {notes}")
        units.append(Unit("\n".join(lines), f"slide {i}"))
    return units, len(prs.slides)


def _table_units(rows: list[list[str]], sheet: str | None) -> list[Unit]:
    """A header-row table as one overview unit plus groups of "column: value" rows, so a chunk
    of rows stays readable without the header line."""
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return []
    header = [h.strip() or f"column {j}" for j, h in enumerate(rows[0], 1)]
    body = rows[1:]
    where = f"{sheet}: " if sheet else ""
    units = [Unit(f"{'Sheet ' + sheet if sheet else 'The table'} has {len(body)} rows and these columns: "
                  f"{', '.join(header)}.", f"{where}columns")]
    budget = thresholds()["s0"]["chunk_words"]
    group: list[str] = []
    words, first = 0, 1
    for i, row in enumerate(body, 1):
        line = "; ".join(f"{h}: {v.strip()}" for h, v in zip(header, row) if v and v.strip())
        n = len(line.split())
        if group and words + n > budget:
            units.append(Unit("\n".join(group), f"{where}rows {first}–{i - 1}"))
            group, words, first = [], 0, i
        group.append(line)
        words += n
    if group:
        units.append(Unit("\n".join(group), f"{where}rows {first}–{len(body)}"))
    return units


def _csv(data: bytes, ext: str) -> tuple[list[Unit], None]:
    text = _decode(data)
    try:
        dialect = csv.excel_tab if ext == ".tsv" else csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return _table_units(list(csv.reader(io.StringIO(text), dialect)), None), None


def _xlsx(data: bytes) -> tuple[list[Unit], int]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    units = []
    for ws in wb.worksheets:
        units += _table_units([["" if v is None else str(v) for v in r] for r in ws.iter_rows(values_only=True)],
                              ws.title)
    return units, len(wb.worksheets)


_HTML_BREAK_RE = re.compile(r"(?i)<br\s*/?>|</(?:p|div|li|tr|h[1-6]|section|article|table)>")
_HTML_DROP_RE = re.compile(r"(?is)<(script|style|head)\b.*?</\1\s*>")
_TAG_RE = re.compile(r"(?s)<[^<>]+>")


def _text(data: bytes, ext: str) -> tuple[list[Unit], None]:
    text = _decode(data)
    if ext in (".html", ".htm", ".xml"):
        text = html.unescape(_TAG_RE.sub(" ", _HTML_BREAK_RE.sub("\n", _HTML_DROP_RE.sub(" ", text))))
    return [Unit(text)], None


def extract(filename: str, data: bytes) -> tuple[list[Unit], int | None]:
    ext = Path(filename).suffix.lower()
    if ext in LEGACY_EXTS:
        raise UnsupportedFile(f"old binary {ext} files are not supported; save it as {LEGACY_EXTS[ext]} and upload again")
    readers = {".pdf": _pdf, ".docx": _docx, ".pptx": _pptx, ".xlsx": _xlsx, ".xlsm": _xlsx}
    try:
        if ext in readers:
            units, pages = readers[ext](data)
        elif ext in (".csv", ".tsv"):
            units, pages = _csv(data, ext)
        elif ext in TEXT_EXTS or not ext:
            units, pages = _text(data, ext)
        else:
            raise UnsupportedFile(f"{ext} files are not supported (try PDF, Word, PowerPoint, Excel, CSV or text)")
    except UnsupportedFile:
        raise
    except Exception as exc:
        raise UnsupportedFile(f"could not read the file ({type(exc).__name__}: {exc})") from exc
    units = [Unit(t, u.location) for u in units if (t := clean_text(u.text))]
    if not units:
        raise UnsupportedFile("no text found (a scanned PDF needs OCR first)")
    return units, pages


def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"(\p{L})-\n(\p{Ll})", r"\1\2", text)       # words hyphenated across PDF lines
    text = re.sub(r"[^\S\n]+", " ", text)
    text = re.sub(r" ?\n ?", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


# ── Chunking ─────────────────────────────────────────────────────────────────
_BOUNDARY_RE = re.compile(r"(?<=[.!?…][\"'”’)\]]*)\s+|\n+")
_WORD_RE = re.compile(r"\S+")


def _pieces(text: str, size: int) -> list[tuple[int, int, int]]:
    """Sentences and lines as (start, end, words); a piece longer than `size` words is cut."""
    out, pos = [], 0
    for m in [*_BOUNDARY_RE.finditer(text), None]:
        end = m.start() if m else len(text)
        words = [w.span() for w in _WORD_RE.finditer(text, pos, end)]
        for k in range(0, len(words), size):
            part = words[k:k + size]
            out.append((part[0][0], part[-1][1], len(part)))
        pos = m.end() if m else end
    return out


def chunk_spans(text: str, size: int, overlap: int) -> list[tuple[int, int]]:
    """Greedy packing of whole pieces up to `size` words; the next chunk starts with the
    trailing pieces of the previous one that fit in `overlap` words."""
    pieces = _pieces(text, size)
    chunks, i = [], 0
    while i < len(pieces):
        j, words = i, 0
        while j < len(pieces) and (j == i or words + pieces[j][2] <= size):
            words += pieces[j][2]
            j += 1
        chunks.append((pieces[i][0], pieces[j - 1][1]))
        if j >= len(pieces):
            break
        k, back = j, 0
        while k - 1 > i and back + pieces[k - 1][2] <= overlap:
            k -= 1
            back += pieces[k][2]
        i = k
    return chunks


# ── Store ────────────────────────────────────────────────────────────────────
@dataclass
class Document:
    info: DocumentInfo
    units: list[Unit]
    chunks: list[tuple[int, int, int]]      # (unit, start, end) offsets into the unit's text
    index: object                           # faiss.Index

    def chunk_text(self, i: int) -> str:
        u, s, e = self.chunks[i]
        return self.units[u].text[s:e]

    def nearest(self, queries: np.ndarray, k: int) -> tuple[np.ndarray, np.ndarray]:
        """FAISS (cosines, chunk ids), one row per query, best first; ids are -1 past the end."""
        return self.index.search(np.ascontiguousarray(queries, dtype="float32"), max(1, min(k, self.index.ntotal)))

    def search(self, queries: np.ndarray, k: int) -> dict[int, float]:
        """Best cosine of each nearby chunk to any of the query vectors."""
        dist, ids = self.nearest(queries, k)
        best: dict[int, float] = {}
        for i, d in zip(ids.ravel().tolist(), dist.ravel().tolist()):
            if i >= 0 and d > best.get(i, -1.0):
                best[i] = d
        return best


def _build_index(vectors: np.ndarray):
    import faiss

    index = faiss.IndexFlatIP(vectors.shape[1])
    index.add(np.ascontiguousarray(vectors, dtype="float32"))
    return index


def _embed_chunks(units: list[Unit], chunks: list[tuple[int, int, int]]) -> np.ndarray:
    return resources.embed([" ".join(units[u].text[s:e].split()) for u, s, e in chunks])


class DocumentStore:
    def __init__(self, root: Path):
        self.root = root
        self._docs: dict[str, Document] = {}
        self._lock = threading.RLock()
        self._loaded = False

    def _ensure_loaded(self) -> None:
        with self._lock:
            if self._loaded:
                return
            self._loaded = True
            for meta in sorted(self.root.glob("*/meta.json")):
                try:
                    self._docs[meta.parent.name] = self._read(meta.parent)
                except Exception as exc:
                    log.warning("skipping stored document %s (%s)", meta.parent.name, exc)

    def _read(self, folder: Path) -> Document:
        import faiss

        meta = json.loads((folder / "meta.json").read_text(encoding="utf-8"))
        info = DocumentInfo(**meta["info"])
        units = [Unit(t, loc) for t, loc in meta["units"]]
        chunks = [tuple(c) for c in meta["chunks"]]
        if info.model != resources.EMBED_MODEL:     # embedder changed: rebuild from the saved chunks
            index = _build_index(_embed_chunks(units, chunks))
            info = info.model_copy(update={"model": resources.EMBED_MODEL})
            doc = Document(info, units, chunks, index)
            self._write(folder, doc)
            return doc
        index = faiss.deserialize_index(np.frombuffer((folder / "index.faiss").read_bytes(), dtype="uint8"))
        return Document(info, units, chunks, index)

    def _write(self, folder: Path, doc: Document) -> None:
        import faiss

        folder.mkdir(parents=True, exist_ok=True)
        # serialise in memory: faiss.write_index cannot open some non-ASCII Windows paths
        (folder / "index.faiss").write_bytes(faiss.serialize_index(doc.index).tobytes())
        (folder / "meta.json").write_text(json.dumps({
            "info": doc.info.model_dump(),
            "units": [[u.text, u.location] for u in doc.units],
            "chunks": [list(c) for c in doc.chunks],
        }, ensure_ascii=False), encoding="utf-8")

    def list(self) -> list[DocumentInfo]:
        self._ensure_loaded()
        with self._lock:
            return sorted((d.info for d in self._docs.values()), key=lambda i: i.created)

    def get(self, ids: list[str]) -> tuple[list[Document], list[str]]:
        self._ensure_loaded()
        with self._lock:
            ids = list(dict.fromkeys(ids))
            return [self._docs[i] for i in ids if i in self._docs], [i for i in ids if i not in self._docs]

    def add(self, filename: str, data: bytes) -> DocumentInfo:
        self._ensure_loaded()
        doc_id = hashlib.sha256(data).hexdigest()[:16]
        with self._lock:
            if doc_id in self._docs:
                return self._docs[doc_id].info
        cfg = thresholds()["s0"]
        units, pages = extract(filename, data)
        chunks = [(u, s, e) for u, unit in enumerate(units)
                  for s, e in chunk_spans(unit.text, cfg["chunk_words"], cfg["chunk_overlap_words"])
                  if e - s >= cfg["min_chunk_chars"] or len(units) == 1]
        if not chunks:
            raise UnsupportedFile("no text found (a scanned PDF needs OCR first)")
        enc = resources.tiktoken_encoding("o200k_base")
        info = DocumentInfo(
            id=doc_id, filename=Path(filename).name, kind=Path(filename).suffix.lower().lstrip(".") or "txt",
            size=len(data), pages=pages, chunks=len(chunks),
            tokens=sum(len(enc.encode(u.text)) for u in units), model=resources.EMBED_MODEL,
            created=datetime.now(timezone.utc).isoformat(timespec="seconds"),
        )
        doc = Document(info, units, chunks, _build_index(_embed_chunks(units, chunks)))
        with self._lock:
            self._write(self.root / doc_id, doc)
            self._docs[doc_id] = doc
        return info

    def delete(self, doc_id: str) -> bool:
        self._ensure_loaded()
        with self._lock:
            if self._docs.pop(doc_id, None) is None:
                return False
            shutil.rmtree(self.root / doc_id, ignore_errors=True)
            return True


@lru_cache(maxsize=None)
def store() -> DocumentStore:
    root = Path(os.environ.get("OPTIMUS_DOC_DIR") or thresholds()["s0"]["store_dir"])
    return DocumentStore(root if root.is_absolute() else BACKEND_DIR / root)
